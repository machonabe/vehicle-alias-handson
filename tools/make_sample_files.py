"""過去の人手承認履歴のサンプルファイル（架空）を data/ に作成する。

使い方: python3 tools/make_sample_files.py   （Excel の作成には openpyxl が必要）
- data/alias_mapping_history.xlsx : ノートブック版で Volume にアップロードして取り込む Excel
    シート「対応履歴」：1行目に表題、2行目に注記、3行目が見出し、4行目以降がデータ（実務の Excel によくある形）
    シート「記入ルール」：列の説明
- data/alias_mapping_history.csv  : SQL エディタ版（read_files）で取り込む CSV（見出し1行＋データ）
"""
import csv
import datetime as dt
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

HEADERS = ["履歴ID", "元システム", "国", "業務領域", "元の名称", "元のコード", "対応づけた共通機種ID",
           "適用開始日", "適用終了日", "状態", "承認部署", "承認日", "元資料"]

D = dt.date
ROWS = [
    ["H001", "CN_SALES_LEGACY", "CN", "SALES", "阿尔法", "CN-X123B", "VEHICLE-001",
     D(2024, 7, 1), D(2025, 5, 31), "APPROVED", "中国営業企画（架空）", D(2024, 6, 20), "Alpha系列 名称対応表 v3.xlsx（架空）"],
    ["H002", "COST_SHEET", "JP", "PRODUCTION", "ALPHA 11 M", None, "VEHICLE-001",
     None, None, "APPROVED", "原価企画（架空）", D(2024, 3, 10), "原価積上げシート_A11.xlsx（架空）"],
    ["H003", "COST_SHEET", "JP", "PRODUCTION", "Ａｌｐｈａ　１１Ｍ", None, "VEHICLE-001",
     None, None, "APPROVED", "原価企画（架空）", D(2024, 3, 10), "原価積上げシート_A11.xlsx（架空）"],
    ["H004", "COST_SHEET", "JP", "SALES", "alpha", None, None,
     None, None, "UNKNOWN", "原価企画（架空）", D(2024, 3, 10), "原価積上げシート_A11.xlsx（架空）"],
    ["H005", "CN_SALES_LEGACY", "CN", "SALES", "阿爾法", None, None,
     None, None, "UNKNOWN", "中国営業企画（架空）", D(2024, 6, 20), "Alpha系列 名称対応表 v3.xlsx（架空）"],
]

RULES = [
    ("履歴ID", "履歴の通し番号"),
    ("元システム", "対応を確認した元のシステム・資料"),
    ("国 / 業務領域", "対象の国と、SALES / PRODUCTION / DEVELOPMENT の区分"),
    ("元の名称 / 元のコード", "元資料に書かれていた名称・コード（表記揺れはそのまま。コードが無い行もある）"),
    ("対応づけた共通機種ID", "人が判断した共通機種ID。分からない場合は空欄"),
    ("適用開始日 / 適用終了日", "対応が有効な期間。分からない場合は空欄"),
    ("状態", "APPROVED（承認済み）/ UNKNOWN（未確認）"),
    ("承認部署 / 承認日", "承認した部署（個人名は書かない）と日付"),
]


def write_csv(path):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        for r in ROWS:
            w.writerow(["" if v is None else (v.isoformat() if isinstance(v, dt.date) else v) for v in r])


def write_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "対応履歴"
    ws["A1"] = "Alpha系列 名称・コード対応履歴（人手で作成・架空データ）"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "各部署が過去に確認した対応をまとめたもの。確定マスタではありません。"
    ws.append(HEADERS)                       # 3行目：見出し
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in ROWS:
        ws.append(r)
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, dt.date):
                c.number_format = "yyyy/mm/dd"
    for col, width in zip("ABCDEFGHIJKLM", [8, 18, 6, 12, 20, 12, 18, 12, 12, 11, 20, 12, 36]):
        ws.column_dimensions[col].width = width

    rules = wb.create_sheet("記入ルール")
    rules.append(["列", "説明"])
    for r in RULES:
        rules.append(list(r))
    wb.save(path)


if __name__ == "__main__":
    DATA.mkdir(exist_ok=True)
    write_csv(DATA / "alias_mapping_history.csv")
    write_xlsx(DATA / "alias_mapping_history.xlsx")
    print("created:", DATA / "alias_mapping_history.csv", DATA / "alias_mapping_history.xlsx")
